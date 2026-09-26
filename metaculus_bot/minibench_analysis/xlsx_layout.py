"""Presentation for recurring reports, following the user's history workbook."""

from __future__ import annotations

import os
from datetime import datetime, timezone

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "17365D"
PALE = "EEF5FB"
INK = "1F2937"


def _header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row, col, label)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def format_report(workbook):
    """Style source sheets and add a formula-linked, keyed competition summary."""
    for ws in workbook:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=10, color=INK)
                cell.alignment = Alignment(vertical="center")
        labels = [c.value for c in ws[1]]
        _header(ws, 1, labels)
        for col, label in enumerate(labels, 1):
            name = str(label or "")
            width = 20
            if name == "minibench":
                width = 42
            elif name in {"title", "question_url", "my_answer_url"}:
                width = 55 if name == "title" else 40
            elif name == "my_prediction":
                width = 34
            elif name == "scoring_methodology":
                width = 115
            ws.column_dimensions[get_column_letter(col)].width = width
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col)
                if name.endswith("brier_skill"):
                    cell.number_format = "0.0%"
                elif name.endswith("_pct"):
                    cell.number_format = '0.0"%"'  # Source values already use 0..100.
                elif "brier" in name or "crps" in name:
                    cell.number_format = "0" if name.endswith("_n") else "0.0000"
                elif name in {"leaderboard_score", "take", "peer_score"} or name.endswith("_peer_avg"):
                    cell.number_format = "#,##0.000"
                if name in {"question_url", "my_answer_url"} and str(cell.value or "").startswith("https://"):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
                if name in {"title", "my_prediction", "scoring_methodology"}:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        if ws.title == "questions":
            for row in range(2, ws.max_row + 1):
                ws.row_dimensions[row].height = 60
        if ws.title == "scoring_notes":
            ws.row_dimensions[2].height = 100

    if "answered" not in workbook.sheetnames or "accuracy" not in workbook.sheetnames:
        return

    def fields(sheet):
        return {c.value: c.column for c in workbook[sheet][1]}

    def rows_by_label(sheet):
        if sheet not in workbook.sheetnames or "minibench" not in fields(sheet):
            return {}
        ws = workbook[sheet]
        return {ws.cell(r, fields(sheet)["minibench"]).value: r for r in range(2, ws.max_row + 1)}

    def ref(sheet, field, row):
        col = fields(sheet).get(field) if sheet in workbook.sheetnames else None
        return f"'{sheet}'!{get_column_letter(col)}{row}" if col and row else None

    def linked(sheet, field, row):
        address = ref(sheet, field, row)
        return f'=IF(ISBLANK({address}),"",{address})' if address else None

    summary = workbook.create_sheet("Summary", 0)
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.tabColor = NAVY
    widths = [44, 12, 12, 14, 16, 12, 14, 16, 12, 18, 10, 12, 18]
    for col, width in enumerate(widths, 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary["A2"] = "Metaculus bot history"
    summary["A2"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    summary["A3"] = "Latest report snapshot. Lower Brier/CRPS is better; higher Brier skill is better."
    summary["A3"].font = Font(name="Arial", size=10, italic=True, color="526579")
    for row in summary.iter_rows(min_row=4, max_row=4, max_col=13):
        for cell in row:
            cell.border = Border(bottom=Side(style="medium", color=NAVY))
    summary.row_dimensions[4].height = 4

    labels = rows_by_label("answered")
    accuracy_rows, ranking_rows = rows_by_label("accuracy"), rows_by_label("ranking")
    cards = [
        (1, 2, "Competitions", len(labels)),
        (3, 4, "Forecasts answered", None),
        (5, 6, "Binary forecasts scored", None),
        (7, 8, "Combined binary Brier", None),
    ]
    for first, last, label, value in cards:
        for r in (6, 7):
            summary.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
            for col in range(first, last + 1):
                c = summary.cell(r, col)
                c.fill = PatternFill("solid", fgColor=PALE)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(name="Arial", size=14 if r == 7 else 10, bold=True, color=NAVY)
        summary.cell(6, first, label)
        summary.cell(7, first, value)
    end = 9 + len(labels)
    summary["C7"] = f"=SUM(B10:B{end})" if labels else 0
    summary["E7"] = f"=SUM(C10:C{end})" if labels else 0
    summary["G7"] = f'=IF(E7=0,"",SUMPRODUCT(C10:C{end},D10:D{end})/E7)' if labels else None
    summary["G7"].number_format = "0.0000"
    summary.row_dimensions[7].height = 28
    _header(
        summary,
        9,
        [
            "Competition",
            "Answered",
            "Binary scored",
            "Binary Brier",
            "Binary skill vs uniform",
            "MC scored",
            "MC Brier",
            "MC skill vs uniform",
            "Numeric scored",
            "Normalized bounded CRPS",
            "Rank",
            "Entries",
            "Leaderboard score",
        ],
    )
    specs = [
        ("answered", "total_answered"),
        ("accuracy", "binary_brier_n"),
        ("accuracy", "binary_brier_mean"),
        ("accuracy", "binary_brier_skill"),
        ("accuracy", "mc_brier_n"),
        ("accuracy", "mc_brier_mean"),
        ("accuracy", "mc_brier_skill"),
        ("accuracy", "numeric_normalized_bounded_crps_n"),
        ("accuracy", "numeric_normalized_bounded_crps_mean"),
        ("ranking", "rank"),
        ("ranking", "leaderboard_entries_returned"),
        ("ranking", "leaderboard_score"),
    ]
    for out_row, (label, source_row) in enumerate(labels.items(), 10):
        summary.cell(out_row, 1, linked("answered", "minibench", source_row))
        indices = {"answered": source_row, "accuracy": accuracy_rows.get(label), "ranking": ranking_rows.get(label)}
        for col, (sheet, field) in enumerate(specs, 2):
            summary.cell(out_row, col, linked(sheet, field, indices[sheet]))
        for cell in summary[out_row]:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(
                vertical="center", horizontal="left" if cell.column == 1 else "right", wrap_text=cell.column == 1
            )
            if out_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FC")
            cell.number_format = (
                "0.0%"
                if cell.column in (5, 8)
                else "0.0000"
                if cell.column in (4, 7, 10)
                else "#,##0.000"
                if cell.column == 13
                else "0"
            )
        summary.row_dimensions[out_row].height = 32
    if labels:
        summary.auto_filter.ref = f"A9:M{end}"
    summary.freeze_panes = "B10"

    note_row = end + 3
    summary.cell(note_row, 1, "Scoring notes").font = Font(name="Arial", size=12, bold=True, color=NAVY)
    notes = [
        ("Binary Brier", "Range 0–1. Mean squared error against the resolved yes/no outcome."),
        ("Multiple-choice Brier", "Range 0–2. Sum of squared errors across all options; not the binary normalization."),
        (
            "Brier skill vs uniform",
            "1 − total Brier / total uniform Brier on the same scored questions. Baseline: binary 0.25; MC 1−1/K. "
            "100% perfect; 0% uniform; negative worse. Higher is better.",
        ),
        (
            "Numeric CRPS",
            "Bounded CDF error divided by the question range. Unknown tails excluded; discrete/date values use interpolation.",
        ),
        (
            "Coverage",
            "Scored counts include only answered, resolved, valid forecasts. Blank scores mean unavailable, not zero.",
        ),
        (
            "Comparison",
            "Compare scores within each metric. Question mix and numeric ranges affect comparisons across competitions.",
        ),
        (
            "Leaderboard",
            "Official rank and score are separate from Brier/CRPS. Entries is the number of leaderboard rows returned.",
        ),
        (
            "Combined binary Brier",
            "Weighted by the number of binary forecasts scored. Forecasts shared across competitions may be counted again.",
        ),
    ]
    for r, (label, text) in enumerate(notes, note_row + 1):
        summary.cell(r, 1, label).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        summary.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
        summary.cell(r, 2, text).font = Font(name="Arial", size=10, color=INK)
        summary.cell(r, 2).alignment = Alignment(vertical="center", wrap_text=True)
        summary.row_dimensions[r].height = 28
    r = note_row + len(notes) + 3
    summary.cell(r, 1, "Generated (UTC)")
    summary.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    summary.cell(r, 2, datetime.now(timezone.utc).replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm"
    repository = os.environ.get("GITHUB_REPOSITORY", "SewerMermaid/metaculus-bot")
    run_id = os.environ.get("GITHUB_RUN_ID")
    summary.cell(r + 1, 1, "Source run" if run_id else "Source")
    summary.cell(r + 1, 2, f"https://github.com/{repository}/actions/runs/{run_id}" if run_id else "Saved report data")
    summary.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=13)
    summary.print_options.horizontalCentered = True
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A3
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.print_area = f"A1:M{r + 1}"
    workbook.active = 0
