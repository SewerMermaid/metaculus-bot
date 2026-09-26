import openpyxl

from metaculus_bot.minibench_analysis.report import write_xlsx


def test_summary_links_match_competition_not_row_position(tmp_path):
    path = tmp_path / "report.xlsx"
    data = {
        "answered": [{"minibench": "A", "total_answered": 2}, {"minibench": "B", "total_answered": 0}],
        "accuracy": [
            {"minibench": "B", "binary_brier_n": 0, "binary_brier_mean": None, "binary_brier_skill": None},
            {"minibench": "A", "binary_brier_n": 2, "binary_brier_mean": 0.0, "binary_brier_skill": 1.0},
        ],
        "ranking": [{"minibench": "A", "rank": 7}],
        "questions": [{"question_url": "https://www.metaculus.com/questions/1/", "brier_score": 0.0}],
    }
    assert write_xlsx(data, str(path))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] == "Summary"
    s = wb["Summary"]
    assert s["D10"].value == "=IF(ISBLANK('accuracy'!C3),\"\",'accuracy'!C3)"
    assert s["D11"].value == "=IF(ISBLANK('accuracy'!C2),\"\",'accuracy'!C2)"
    assert s["K11"].value is None
    assert s["E10"].value == "=IF(ISBLANK('accuracy'!D3),\"\",'accuracy'!D3)"
    assert s["E10"].number_format == "0.0%"
    assert s["E9"].value == "Binary skill vs uniform"
    assert s["H9"].value == "MC skill vs uniform"
    assert s["G7"].value == '=IF(E7=0,"",SUMPRODUCT(C10:C11,D10:D11)/E7)'
    assert wb["accuracy"]["C3"].value == 0
    assert wb["accuracy"]["C2"].value is None
    assert wb["questions"]["A2"].hyperlink.target.endswith("/1/")
    assert wb["questions"].freeze_panes == "B2"
    assert s["D10"].number_format == "0.0000"


def test_empty_report_does_not_create_reversed_formula_ranges(tmp_path):
    path = tmp_path / "empty.xlsx"
    assert write_xlsx({"answered": [], "accuracy": []}, str(path))
    wb = openpyxl.load_workbook(path)
    assert wb["Summary"]["C7"].value == 0
    assert wb["Summary"]["G7"].value is None
